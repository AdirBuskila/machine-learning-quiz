# -*- coding: utf-8 -*-
"""Dump raw text from every ML exam / solutions / practice file to tools/raw/.

Hebrew RTL PDFs mangle differently under pdfplumber vs PyMuPDF(fitz), so we dump
BOTH per PDF and let the structuring step pick the cleaner one. DOCX via python-docx.
The big ZIP of exams is unzipped and included. Emits tools/raw/manifest.json with
per-file metadata (pages, image counts per page, char counts, short previews).

Run:  PYTHONUTF8=1 python tools/extract_exams.py
Prints ASCII only (Windows console is cp1255); Hebrew goes to UTF-8 files.
"""
import os, re, json, zipfile, pathlib, traceback
import pdfplumber
import fitz  # PyMuPDF
import docx  # python-docx

SRC = pathlib.Path(r"C:\Users\Adir\Desktop\BSC\שנה ב\סמסטר ג\למידת מכונה")
HERE = pathlib.Path(__file__).parent
RAW = HERE / "raw"
TXT = RAW / "txt"
ZIPX = RAW / "zip_extracted"
RAW.mkdir(parents=True, exist_ok=True)
TXT.mkdir(parents=True, exist_ok=True)

MAX_BYTES = 9 * 1024 * 1024          # skip huge aggregates (Daniel 45MB, big summaries)
# folders to harvest exam/practice files from (relative to SRC)
SCAN_DIRS = ["מבחנים", "קבצי תרגול למבחנים"]
# explicit skips (huge collections handled separately / not individual exams)
SKIP_NAME_SUBSTR = ["מבחנים דניאל", "קובץ עם מלא מבחנים.zip"]


def plumber_text(path):
    out = []
    with pdfplumber.open(path) as doc:
        for i, page in enumerate(doc.pages, 1):
            out.append(f"\n----PAGE {i}----\n")
            out.append(page.extract_text() or "")
    return "".join(out)


def fitz_dump(path):
    """Return (text, [images_per_page])."""
    out, imgs = [], []
    doc = fitz.open(path)
    for i, page in enumerate(doc, 1):
        out.append(f"\n----PAGE {i}----\n")
        out.append(page.get_text("text"))
        try:
            imgs.append(len(page.get_images(full=True)))
        except Exception:
            imgs.append(0)
    doc.close()
    return "".join(out), imgs


def docx_text(path):
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    # include tables (answer keys often live in tables)
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def slug(idx):
    return f"{idx:03d}"


def gather_files():
    """Yield (abs_path, rel_path, origin) for every candidate file."""
    seen = set()
    for d in SCAN_DIRS:
        base = SRC / d
        if not base.exists():
            continue
        for p in sorted(base.rglob("*")):
            if not p.is_file():
                continue
            if p.suffix.lower() not in (".pdf", ".docx"):
                continue
            if any(s in p.name for s in SKIP_NAME_SUBSTR):
                continue
            if p.stat().st_size > MAX_BYTES:
                continue
            key = p.resolve()
            if key in seen:
                continue
            seen.add(key)
            yield p, str(p.relative_to(SRC)), "loose"


def gather_zip():
    """Unzip the big archive and yield its pdf/docx files."""
    zpath = SRC / "מבחנים" / "מבחן לדוגמה" / "קובץ עם מלא מבחנים.zip"
    if not zpath.exists():
        return
    ZIPX.mkdir(parents=True, exist_ok=True)
    z = zipfile.ZipFile(zpath)
    for info in z.infolist():
        if info.is_dir():
            continue
        name = info.filename
        if not (info.flag_bits & 0x800):
            for enc in ("utf-8", "cp862", "cp1255"):
                try:
                    name = info.filename.encode("cp437").decode(enc)
                    break
                except Exception:
                    pass
        if not name.lower().endswith((".pdf", ".docx")):
            continue
        safe = re.sub(r"[^A-Za-z0-9._-]", "_", name)
        dest = ZIPX / safe
        try:
            with z.open(info) as src, open(dest, "wb") as f:
                f.write(src.read())
        except Exception:
            continue
        yield dest, "ZIP/" + name, "zip"


def main():
    manifest = []
    idx = 0
    files = list(gather_files()) + list(gather_zip())
    for path, rel, origin in files:
        idx += 1
        st = slug(idx)
        rec = {"idx": idx, "stem": st, "rel": rel, "origin": origin,
               "fmt": path.suffix.lower().lstrip("."), "bytes": path.stat().st_size,
               "pages": None, "img_per_page": [], "plumber_chars": 0, "fitz_chars": 0,
               "plumber_preview": "", "fitz_preview": "", "error": None}
        try:
            if path.suffix.lower() == ".pdf":
                try:
                    pt = plumber_text(str(path))
                except Exception as e:
                    pt = ""
                    rec["error"] = f"plumber:{e.__class__.__name__}"
                ft, imgs = fitz_dump(str(path))
                (TXT / f"{st}.plumber.txt").write_text(pt, encoding="utf-8")
                (TXT / f"{st}.fitz.txt").write_text(ft, encoding="utf-8")
                rec.update(pages=len(imgs), img_per_page=imgs,
                           plumber_chars=len(pt), fitz_chars=len(ft),
                           plumber_preview=pt[:600], fitz_preview=ft[:600])
            else:  # docx
                dt = docx_text(path)
                (TXT / f"{st}.docx.txt").write_text(dt, encoding="utf-8")
                rec.update(pages=None, plumber_chars=len(dt), fitz_chars=len(dt),
                           plumber_preview=dt[:600], fitz_preview=dt[:600])
        except Exception as e:
            rec["error"] = f"{e.__class__.__name__}: {e}"
            (RAW / "extract_errors.log").open("a", encoding="utf-8").write(
                f"{rel}\n{traceback.format_exc()}\n")
        manifest.append(rec)
        print(f"[{st}] {origin:5s} {rec['fmt']:4s} pages={rec['pages']} "
              f"pl={rec['plumber_chars']} fz={rec['fitz_chars']} "
              f"imgs={sum(rec['img_per_page'])} err={rec['error']}")

    (RAW / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"\nDONE. {len(manifest)} files. manifest -> {RAW/'manifest.json'}")


if __name__ == "__main__":
    main()
